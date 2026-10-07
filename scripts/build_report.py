"""Build the report from measured output/verification.json and screenshots."""
import csv
import json
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
stats = json.loads((OUT / "verification.json").read_text(encoding="utf-8"))
rows = list(csv.DictReader((OUT / "telemetry_raw.csv").open(encoding="utf-8-sig")))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(6)
for name, size in [("Title", 23), ("Subtitle", 13), ("Heading 1", 15), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(6)
    if name.startswith("Heading"):
        style.font.bold = True
doc.core_properties.author = "Lê Văn Duy"
doc.core_properties.title = "Báo cáo bài thực hành số 2 IoT"
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    if style.type == 1:
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if attr.endswith("Theme"):
                del fonts.attrib[attr]
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")


def p(text, style=None):
    return doc.add_paragraph(text, style)


def heading(text):
    doc.add_heading(text, level=1)


def table(headers, data):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    if len(headers) == 2:
        t.autofit = False
        t.columns[0].width = Inches(1.65)
        t.columns[1].width = Inches(5.25)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E8EEEE")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in data:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for row in t.rows:
        if len(headers) == 2:
            row.cells[0].width = Inches(1.65)
            row.cells[1].width = Inches(5.25)
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
                for run in para.runs:
                    run.font.size = Pt(10.5)
    p("")
    return t


def caption(text):
    para = p(text)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:
        run.italic = True
        run.font.size = Pt(10)


def page():
    doc.add_page_break()


def architecture():
    im = Image.new("RGB", (1440, 430), "white")
    draw = ImageDraw.Draw(im)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 27)
    small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 22)
    blocks = [(20, 40, 325, 155, "ESP32 Wokwi", "DHT22 + HC-SR04"),
              (390, 40, 695, 155, "aMQTT + collector", "Python / durable inbox"),
              (760, 40, 1065, 155, "MongoDB Time-Series", "sensor_raw / 7 days"),
              (760, 270, 1065, 385, "Python + pandas", "clean / resample / features"),
              (390, 270, 695, 385, "MongoDB", "sensor_processed"),
              (20, 270, 325, 385, "Streamlit", "dashboard + CSV")]
    for x1,y1,x2,y2,title,sub in blocks:
        draw.rounded_rectangle((x1,y1,x2,y2), radius=8, fill="#edf4f3", outline="#326f65", width=3)
        draw.text(((x1+x2)/2,y1+28),title,font=font,fill="#183b35",anchor="mm")
        draw.text(((x1+x2)/2,y1+77),sub,font=small,fill="#34464b",anchor="mm")
    for start,end in [((325,97),(390,97)),((695,97),(760,97)),((912,155),(912,270)),((760,327),(695,327)),((390,327),(325,327))]:
        draw.line((start,end), fill="#606060",width=4)
        x,y=end
        if start[1] != end[1]: tri=[(x,y),(x-8,y-14),(x+8,y-14)]
        elif start[0] < x: tri=[(x,y),(x-14,y-8),(x-14,y+8)]
        else: tri=[(x,y),(x+14,y-8),(x+14,y+8)]
        draw.polygon(tri,fill="#606060")
    draw.text((1110,55),"Test input",font=font,fill="#934b28")
    draw.text((1110,97),"demo.py",font=small,fill="#934b28")
    draw.text((1110,132),"via MQTT",font=small,fill="#934b28")
    path=OUT / "architecture.png"
    im.save(path)
    return path


p("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG").alignment = WD_ALIGN_PARAGRAPH.CENTER
p("BÁO CÁO BÀI THỰC HÀNH SỐ 2", "Title")
p("Thu thập lưu trữ và tiền xử lý dữ liệu IoT", "Subtitle")
p("Học phần IoT và Ứng dụng INT14149\nSinh viên Lê Văn Duy    |    Mã sinh viên B23DCAT074")
heading("1 Mục tiêu và phạm vi triển khai")
p("Bài thực hành xây dựng luồng dữ liệu từ cảm biến đến cơ sở dữ liệu và ứng dụng giám sát. Hệ thống tiếp nhận JSON qua MQTT, lưu dữ liệu chuỗi thời gian, xử lý giá trị thiếu và bất thường, sau đó hiển thị kết quả bằng dashboard. Phương án triển khai chạy trực tiếp trên Windows bằng Python, không sử dụng Docker.")
p("Firmware ESP32 đã biên dịch thành công cho Wokwi. Kết quả thực nghiệm trong báo cáo được đo bằng bộ phát kiểm thử Python qua MQTT và MongoDB thật trên máy tính; không được xem là số đo cảm biến ESP32. Việc chạy tương tác Wokwi và đo độ trễ thiết bị là bước xác nhận bổ sung khi demo.")
doc.add_picture(str(architecture()), width=Inches(6.8))
caption("Hình 1 Kiến trúc hệ thống và vị trí bộ phát kiểm thử")
table(["Thành phần", "Lựa chọn"], [
    ["Thiết bị và cảm biến", "ESP32, DHT22, HC-SR04 trên Wokwi"],
    ["MQTT và ngôn ngữ", "aMQTT 0.11.3, paho-mqtt 2.1.0, Python 3.12"],
    ["Lưu trữ và xử lý", f"MongoDB {stats['mongo_version']}, PyMongo, pandas, NumPy"],
    ["Ứng dụng", "Streamlit 1.50.0 và Plotly; http://127.0.0.1:8502"]])

page()
heading("2 Thu thập và thiết kế lưu trữ")
p("ESP32 sử dụng chân DHT22 GPIO21, TRIG GPIO5, ECHO GPIO18 và LED GPIO23. Mỗi 5 giây thiết bị tạo một bản tin JSON trên topic ptit/B23DCAT074/sensors. Thời gian được đồng bộ NTP trước khi gửi; mỗi lần khởi động có boot_id mới để phân biệt việc sequence quay về đầu.")
table(["Nhóm", "Trường và ý nghĩa"], [
    ["Nhận dạng", "device_id, boot_id, source; source là wokwi hoặc test"],
    ["Thời gian", "timestamp_ms là UTC epoch ms; uptime_s là thời gian hoạt động"],
    ["Quan sát", "temperature (°C), humidity (%RH), distance_cm (cm), rssi (dBm)"],
    ["Toàn vẹn", "sequence tăng theo lần lấy mẫu; event_id được tính từ thiết bị, phiên và sequence"]])
p("Collector kiểm tra JSON, kiểu dữ liệu, thời gian và thông tin nhận dạng. Một giá trị cảm biến thiếu có thể là null; chuỗi NaN, số vô hạn và bản tin thiếu toàn bộ cảm biến bị loại. Số hữu hạn ngoài miền vẫn được lưu ở dữ liệu gốc để giữ bằng chứng bất thường cho bước xử lý.")
table(["Nơi lưu", "Thiết kế và mục đích"], [
    ["sensor_raw", "MongoDB Time-Series: timeField=timestamp, metaField=meta, granularity=seconds. Lưu cảm biến và thời điểm nhận."],
    ["meta", "device_id, boot_id, source. Chỉ mục event_id hỗ trợ tìm bản tin đã ghi."],
    ["sensor_processed", "Collection thường, upsert theo thiết bị, phiên và mốc cửa sổ. Cho phép cập nhật kết quả khi có thêm dữ liệu."],
    ["SQLite inbox", "Bộ đệm cục bộ bền vững, khóa chính event_id; ghi nhận thời điểm nhận, hoàn tất ghi và lỗi."],
    ["Retention", "Raw và processed giữ 604800 giây (7 ngày). TTL dọn bất đồng bộ; không xóa chính xác ngay giây hết hạn."]])
p("MongoDB Time-Series không có unique index. Vì vậy chỉ chạy một collector, nhận diện trùng tại inbox và kiểm tra event_id trước khi thử ghi lại. Nếu MongoDB tạm dừng, bản tin còn ở inbox để tiếp tục ghi khi khôi phục. Các truy vấn phân tích cũng loại trùng theo event_id. Thiết kế này phù hợp bài thực hành cục bộ nhưng không phải bảo đảm exactly-once phân tán.")
p("PubSubClient trên ESP32 gửi QoS 0; bộ phát Python dùng QoS 1. Collector xác nhận bản tin QoS 1 sau khi commit vào inbox. Broker aMQTT giữ trạng thái trong RAM, nên chưa bảo đảm hàng đợi tồn tại qua lần khởi động lại broker.")

page()
heading("3 Quy trình tiền xử lý")
p("Chương trình preprocess.py đọc dữ liệu 60 phút gần nhất và chạy lại mỗi 5 giây. Dữ liệu được chia theo device_id, boot_id và source để không nội suy xuyên lần khởi động lại hay trộn dữ liệu thử với dữ liệu thiết bị. Dữ liệu gốc được giữ nguyên.")
table(["Bước", "Cách thực hiện và giới hạn"], [
    ["Kiểm tra miền", "Nhiệt độ −40 đến 80 °C; độ ẩm 0 đến 100 %RH; khoảng cách 2 đến 400 cm. Ngoài miền được gắn cờ và đổi thành giá trị thiếu ở nhánh xử lý."],
    ["Phát hiện IQR", "Với ít nhất 8 mẫu hợp lệ: IQR = Q3 − Q1. Đánh dấu ngoài [Q1 − 1.5 IQR; Q3 + 1.5 IQR]. Bỏ kiểm tra thống kê nếu IQR bằng 0."],
    ["Resampling", "Gom vào cửa sổ 10 giây theo UTC và lấy trung bình; lưu sample_count cùng số ô thiếu và bất thường."],
    ["Xử lý thiếu", "Nội suy tuyến tính theo thời gian, chỉ với khoảng thiếu bên trong dài tối đa 2 cửa sổ. Không điền đầu/cuối hoặc khoảng thiếu dài."],
    ["Tạo đặc trưng", "Rolling mean trên 3 cửa sổ; delta là chênh lệch hai cửa sổ liên tiếp. Không tính delta qua giá trị còn thiếu."],
    ["Chuẩn hóa", "Min-Max theo miền vật lý cố định: (x − min) / (max − min). Các nguồn dùng cùng thang chuẩn hóa."],
    ["Ghi kết quả", "Upsert vào sensor_processed, lưu processed_at, window_s và cờ imputed để truy vết."]])
p("Lựa chọn miền vật lý cố định tránh việc cùng một giá trị có điểm chuẩn hóa khác nhau sau mỗi lần làm mới. Nội suy chỉ áp dụng cho khoảng thiếu ngắn nhằm hạn chế việc tạo ra đường biểu diễn trông liên tục trong khi thiết bị đã mất kết nối lâu.")
p("IQR được tính trên dữ liệu của từng phiên trong khoảng truy vấn, vì vậy thay đổi môi trường thật có thể bị đánh dấu nhầm. Cần đối chiếu raw với processed và điều chỉnh ngưỡng theo bài toán. Kết quả ở đầu/cuối cửa sổ truy vấn có thể thay đổi khi dữ liệu mới đến; các cờ chất lượng không được bỏ khi xuất dữ liệu.")
heading("4 Thiết kế phép thử")
p(f"Bộ phát Python gửi {stats['nominal_count']} vị trí sequence với chu kỳ {stats['interval_s']:g} giây trong phép thử tăng tốc. Cố ý bỏ sequence 22, lặp sequence 10, để thiếu độ ẩm ở 12, đặt nhiệt độ 78 °C ở 18 và khoảng cách 900 cm ở 30. Ba bản tin sai cấu trúc hoặc kiểu được gửi thêm để kiểm tra validation. Chu kỳ firmware ESP32 vẫn là 5 giây.")

page()
heading("5 Kết quả thực nghiệm")
when = datetime.fromtimestamp(stats["started"], ZoneInfo("Asia/Ho_Chi_Minh")).strftime("%d/%m/%Y %H:%M:%S")
p(f"Phiên kiểm thử bắt đầu lúc {when} (UTC+7), boot_id {stats['boot_id']}. Số liệu dưới đây được truy vấn lại từ MongoDB và sổ nhận sau khi bộ phát kết thúc, không phải giá trị dự kiến điền sẵn.")
table(["Chỉ tiêu", "Kết quả đo", "Nhận xét"], [
    ["Bản ghi duy nhất", stats["unique_received"], "Bằng số bản tin hợp lệ đã gửi"],
    ["Bản tin trùng bị chặn", stats["events"].get("duplicate", 0), "Không nhân đôi raw"],
    ["Bản tin bị từ chối", stats["events"].get("rejected", 0), "JSON/kiểu/nhận dạng không hợp lệ"],
    ["Sequence thiếu ở giữa", stats["sequence_missing"], "Đúng vị trí 22 đã chủ động bỏ"],
    ["Ô cảm biến thiếu", stats["missing_cells"], "humidity ở sequence 12"],
    ["Ô bị gắn cờ bất thường", stats["outlier_cells"], "Kiểm tra IQR và miền vật lý"],
    ["Cửa sổ sau xử lý", stats["processed_windows"], "Cửa sổ 10 giây"]])
caption("Bảng 1 Kết quả kiểm thử qua MQTT và MongoDB cục bộ")
sample = []
for r in rows[:5]:
    sample.append([r["sequence"], f"{float(r['temperature']):.2f}", f"{float(r['humidity']):.2f}", f"{float(r['distance_cm']):.2f}"])
table(["Sequence", "Nhiệt độ °C", "Độ ẩm %RH", "Khoảng cách cm"], sample)
caption("Bảng 2 Năm bản ghi đầu của nguồn kiểm thử Python")
p("Độ trễ nhận đến lưu được tính từ khi callback MQTT nhận bản tin tới khi MongoDB xác nhận ghi. Độ trễ publisher đến lưu tính từ timestamp trong bản tin; trong phép thử này hai chương trình dùng cùng đồng hồ máy tính.")
table(["Độ trễ", "Trung bình", "Phân vị 95"], [
    ["Nhận MQTT đến lưu DB", f"{stats['receive_to_db_mean_ms']:.2f} ms", f"{stats['receive_to_db_p95_ms']:.2f} ms"],
    ["Python publisher đến lưu DB", f"{stats['publisher_to_db_mean_ms']:.2f} ms", f"{stats['publisher_to_db_p95_ms']:.2f} ms"]])
p("Đây là kết quả của một phiên ngắn trên máy cục bộ, chưa đại diện cho mạng Internet hay hiệu năng nhiều thiết bị. Độ trễ tới màn hình còn thêm thời gian truy vấn và chu kỳ làm mới dashboard; chưa đo thời điểm trình duyệt vẽ xong từng mẫu.")

page()
heading("6 Ứng dụng giám sát")
p("Dashboard có bộ lọc nguồn và khoảng thời gian, thẻ giá trị hiện tại, biểu đồ cảm biến, dữ liệu sau xử lý, cờ chất lượng và chức năng xuất CSV. Nguồn kiểm thử hiển thị cảnh báo riêng, tránh nhầm với dữ liệu từ ESP32. Các cảm biến dùng trục riêng để không trộn đơn vị.")
doc.add_picture(str(OUT / "dashboard.png"), width=Inches(6.8))
caption("Hình 2 Dashboard chạy trên trình duyệt với dữ liệu kiểm thử thực tế")
heading("7 Lưu trữ và hiệu năng")
p(f"Thống kê MongoDB tại thời điểm kiểm chứng ghi nhận dung lượng logic collection raw là {stats['db_logical_bytes']} byte và dung lượng cấp phát lưu trữ là {stats['db_storage_bytes']} byte. Chỉ số này có overhead bucket và cấp phát, không phải toàn bộ thư mục database hay kích thước trung bình ổn định cho mỗi bản ghi.")
p("Với chu kỳ 5 giây, một thiết bị tạo tối đa 17280 mẫu/ngày và 120960 mẫu/7 ngày nếu gửi liên tục. Đây là ước tính theo chu kỳ, không phải số lượng đã đo. Cần chạy dài hơn và tăng số thiết bị để đánh giá tốc độ ghi, bộ nhớ, dung lượng thực và độ trễ phân vị cao.")
p("Giới hạn thời gian lưu kiểm soát tăng trưởng database, nhưng inbox đang chờ ghi vẫn có thể tăng khi MongoDB lỗi lâu. Hướng cải tiến là giới hạn dung lượng bộ đệm, cảnh báo khi tồn đọng, chia lô ghi và dùng broker có cơ chế lưu hàng đợi bền vững.")

page()
heading("8 Đánh giá và hướng dẫn demo")
table(["Hạng mục", "Trạng thái và bằng chứng"], [
    ["Firmware ESP32", "PlatformIO build thành công. RAM 45204 byte, flash 765705 byte. Chưa xác nhận chạy tương tác Wokwi trong phép thử này."],
    ["Thu thập và lưu trữ", "MQTT broker thật, collector và MongoDB Time-Series đã chạy; verification.json ghi số đo và retention."],
    ["Tiền xử lý", "Đã chạy trên raw, có kết quả CSV và cờ chất lượng. 15 kiểm thử tự động đạt, bao phủ validation, trùng, lỗi ghi, nội suy và outlier."],
    ["Dashboard", "Đã mở trình duyệt, hiển thị biểu đồ và xuất dữ liệu. Hình 2 là ảnh chụp ứng dụng."],
    ["Mã nguồn và tài liệu", "Có firmware, các script Python, cấu hình, requirements và README. Cần đưa bộ mã vào GitHub của sinh viên khi nộp."]])
p("Để demo, mở thư mục iot_lab2 trong VS Code và chạy .venv\\Scripts\\python.exe run.py. Mở http://127.0.0.1:8502. Trong terminal thứ hai chạy .venv\\Scripts\\python.exe demo.py để phát mẫu mỗi 5 giây; chọn nguồn Python kiểm thử. Chạy verify_run.py sau khi kết thúc để xuất số liệu mới.")
p("Để xác nhận phần thiết bị, mở thư mục firmware trong VS Code, Build rồi Wokwi: Start Simulator. Chọn nguồn ESP32 Wokwi trên dashboard và thay đổi các cảm biến. Wokwi phải dùng gateway trong VS Code để truy cập host.wokwi.internal:1884. Chụp Serial và dashboard để bổ sung bằng chứng thiết bị trước khi nộp.")
heading("9 Khó khăn và cách giải quyết")
p("Yêu cầu không dùng Docker được giải quyết bằng bộ khởi động Python và MongoDB portable. Việc mất kết nối database được xử lý bằng inbox bền vững. Vấn đề sequence reset được tách bằng boot_id. Hạn chế unique index của Time-Series được xử lý bằng collector đơn và kiểm tra event_id khi ghi lại. Token bài 1 không được đưa vào firmware bài 2.")
p("Pipeline Python đã được kiểm chứng với dữ liệu thử có kiểm soát. Phần cần hoàn thiện khi demo là chạy cảm biến Wokwi thực tế, xác nhận đồng bộ thời gian và đánh giá phục hồi mạng. Khi mở rộng khỏi máy cục bộ cần MQTT/TLS, xác thực database và cơ chế hàng đợi bền vững.")
heading("Tài liệu tham khảo")
for text in ["1. Đề Bài thực hành số 2 IoT, học phần INT14149, mục 3 đến 7.",
             "2. MongoDB Documentation: mongodb.com/docs/manual/core/timeseries-collections/ và timeseries/timeseries-limitations/.",
             "3. aMQTT Documentation: amqtt.io; Eclipse Paho Python: eclipse.dev/paho/clients/python/docs/.",
             "4. Wokwi VS Code Project Configuration: docs.wokwi.com/vscode/project-config."]:
    para=p(text)
    for run in para.runs:
        run.font.size=Pt(10)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("B23DCAT074  |  ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
path = OUT / "Bao_cao_Bai2_Le_Van_Duy_B23DCAT074.docx"
doc.save(path)
print(path)
